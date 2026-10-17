"""Read a training agenda workbook (the MT_SLPHC_Master_Training_Agenda layout) into agenda rows.

Layout: a "DAYn - title" line starts each day; session rows carry TIME ("08:30 - 10:00"), a duration,
the ACTIVITY and the RESPONSIBLE PERSON(S); rows without a time continue the previous session (the
chapter sections it covers). Breakfast, registration, lunch and attendance rows are not sessions a
facilitator reports on, so they are left out. CSV files with the columns day, topic, planned_start,
planned_end, details, responsible are read as they are."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook

DAY_RE = re.compile(r"^\s*DAY\s*(\d+)\s*[-–—:]?\s*(.*)$", re.IGNORECASE)
DAY_RANGE_RE = re.compile(r"^\s*DAY\s*\d+\s*[-–—]\s*DAY\s*\d+", re.IGNORECASE)  # "DAY 1 - DAY 5": the title of a block, not a day
NTH_DAY_RE = re.compile(r"\((\d+)\s*(?:st|nd|rd|th)\s+day\)", re.IGNORECASE)  # "DAY 1 (6th Day)": a parallel track numbered from 1, 6th day overall
END_RE = re.compile(r"^\s*SUMMARY OF WORK HOURS", re.IGNORECASE)
TIME_RE = re.compile(r"^\s*(\d{1,2})[:.](\d{2})\s*[-–—]+\s*(\d{1,2})[:.](\d{2})\s*$")
SKIP_RE = re.compile(r"^(registration|attendance|lunch|breakfast|tea break|coffee break|break)\b", re.IGNORECASE)
BULLET_RE = re.compile(r"^[\s·•\-–*]+")


@dataclass
class AgendaRow:
    day: int
    topic: str
    planned_start: str | None = None
    planned_end: str | None = None
    details: list[str] = field(default_factory=list)
    responsible: str | None = None

    def as_dict(self) -> dict:
        return {"day": self.day, "topic": self.topic, "planned_start": self.planned_start, "planned_end": self.planned_end,
                "details": " / ".join(d for d in self.details if d) or None, "responsible": self.responsible}


@dataclass
class ParsedAgenda:
    rows: list[AgendaRow]
    day_titles: dict[int, str]
    warnings: list[str]


def _clean(text) -> str:
    s = "" if text is None else str(text)
    s = BULLET_RE.sub("", s.replace("\xa0", " ")).strip()
    return re.sub(r"\s+", " ", s)


def _time(text) -> tuple[str | None, str | None]:
    m = TIME_RE.match("" if text is None else str(text))
    if not m:
        return None, None
    h1, m1, h2, m2 = (int(x) for x in m.groups())
    if h1 > 23 or h2 > 23:
        return None, None
    return f"{h1:02d}:{m1:02d}", f"{h2:02d}:{m2:02d}"


def parse_rows(rows: list[list], skip_breaks: bool = True) -> ParsedAgenda:
    """rows: cell values per spreadsheet row (TIME, DURATION, ACTIVITY, RESPONSIBLE in the first columns
    the layout uses: A time, B day hours, C duration, D activity, E responsible)."""
    out: list[AgendaRow] = []
    titles: dict[int, str] = {}
    warnings: list[str] = []
    day = 0
    current: AgendaRow | None = None
    for cells in rows:
        cells = list(cells) + [None] * 6
        first = _clean(cells[0])
        activity_cell = _clean(cells[3])
        if END_RE.match(first):
            break
        # a day heading can sit in column A or D; "DAY 1 - DAY 5" titles a block and is skipped;
        # "DAY 1 (6th Day)" numbers a parallel track from 1 and is filed under its overall day (6)
        for candidate in (first, activity_cell):
            if DAY_RANGE_RE.match(candidate):
                break
            m = DAY_RE.match(candidate)
            if m and not TIME_RE.match(first):
                nth = NTH_DAY_RE.search(candidate)
                n = int(nth.group(1)) if nth else int(m.group(1))
                if first or day > 0:  # a heading row in the body (not the table of contents at the top)
                    day = n
                    current = None
                    titles.setdefault(n, _clean(NTH_DAY_RE.sub("", m.group(2))) or f"Day {n}")
                else:
                    titles.setdefault(n, _clean(m.group(2)) or f"Day {n}")
                break
        else:
            if day == 0:
                continue
            start, end = _time(cells[0])
            duration = _clean(cells[2])
            responsible = _clean(cells[4]) or None
            if not activity_cell and not start:
                continue
            if activity_cell.upper() in ("ACTIVITY",) or first.upper() == "TIME":
                continue
            is_session = bool(start) or bool(duration)
            looks_like_detail = bool(re.match(r"^\d+(\.\d+)+", activity_cell)) or activity_cell[:1].islower()
            if start and current is not None and current.planned_start is None and looks_like_detail:
                # the time came one row below the session title (Day 1, Chapter 3)
                current.planned_start, current.planned_end = start, end
                current.details.append(activity_cell)
                continue
            if is_session and activity_cell:
                if skip_breaks and SKIP_RE.match(activity_cell):
                    current = None
                    continue
                current = AgendaRow(day=day, topic=activity_cell.rstrip(":").strip(), planned_start=start, planned_end=end, responsible=responsible)
                out.append(current)
            elif activity_cell and current is not None:
                current.details.append(activity_cell)
            elif activity_cell:
                warnings.append(f"Day {day}: '{activity_cell[:60]}' has no session to attach to")
    for r in out:
        if not r.planned_start:
            warnings.append(f"Day {r.day}: '{r.topic[:60]}' has no planned time")
    return ParsedAgenda(rows=out, day_titles=titles, warnings=warnings)


def parse_xlsx(data: bytes, skip_breaks: bool = True) -> ParsedAgenda:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    return parse_rows([list(r) for r in ws.iter_rows(values_only=True)], skip_breaks)


def parse_csv(data: bytes) -> ParsedAgenda:
    """A plain list: day, topic, planned_start, planned_end[, details[, responsible]] with a header row."""
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    out: list[AgendaRow] = []
    warnings: list[str] = []
    keys = {k.strip().lower(): k for k in (reader.fieldnames or [])}

    def get(row, *names):
        for n in names:
            if n in keys:
                return _clean(row.get(keys[n]))
        return ""

    for i, row in enumerate(reader, start=2):
        day = get(row, "day", "training day")
        topic = get(row, "topic", "session", "session / topic", "activity")
        if not day.isdigit() or not topic:
            warnings.append(f"Line {i}: needs a day number and a topic")
            continue
        s = _time(f"{get(row, 'planned_start', 'start')} - {get(row, 'planned_end', 'end')}")
        out.append(AgendaRow(day=int(day), topic=topic, planned_start=s[0], planned_end=s[1], details=[get(row, "details", "content")], responsible=get(row, "responsible") or None))
    return ParsedAgenda(rows=out, day_titles={}, warnings=warnings)


def parse_file(filename: str, data: bytes, skip_breaks: bool = True) -> ParsedAgenda:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data, skip_breaks)
    return parse_csv(data)
