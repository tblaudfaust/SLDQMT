"""Standard reports as XLSX or PDF. Each report is a title plus one or more
(sheet name, headers, rows) tables built from the dashboard queries, so the
numbers always match the dashboard for the same filters."""

import io
from dataclasses import dataclass
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import District, ErrorCategory, ErrorRecord, ErrorStatus, Team, User
from app.schemas.dashboard import Filters
from app.services import dashboard_service as ds

REPORT_KINDS = {
    "daily_summary": "Daily error follow-up summary",
    "district_register": "District error register",
    "overdue": "Overdue follow-ups",
    "category_analysis": "Error category analysis",
    "team_performance": "Team performance",
    "monitor_activity": "Field Monitor activity",
    "error_detail": "Error detail sheet",
    "full_export": "Full export",
}


@dataclass
class Sheet:
    name: str
    headers: list[str]
    rows: list[list]
    landscape: bool = False


@dataclass
class Report:
    kind: str
    title: str
    sheets: list[Sheet]
    generated_at: datetime
    filters: str


def _fmt(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if v is None:
        return ""
    return v


def _filters_text(f: Filters) -> str:
    parts = [f"{k}={v}" for k, v in f.model_dump(exclude_none=True).items() if v not in (False, None, [])]
    return "; ".join(parts) if parts else "none"


def build(db: Session, user: User, kind: str, f: Filters, error_id: str | None = None) -> Report:
    if kind not in REPORT_KINDS:
        raise ValueError(f"Unknown report kind {kind}")
    now = datetime.now(timezone.utc)
    sheets: list[Sheet] = []

    if kind == "daily_summary":
        s = ds.summary(db, user, f)
        sheets.append(
            Sheet(
                "Summary",
                ["Measure", "Value"],
                [
                    ["Total errors", s.total],
                    ["Resolved", s.resolved],
                    ["Unresolved", s.unresolved],
                    ["Overdue follow-ups", s.overdue],
                    ["Resolution rate (%)", s.resolution_rate],
                    ["Median days to resolve", s.median_days_to_resolve if s.median_days_to_resolve is not None else ""],
                    ["Last tablet sync", _fmt(s.last_sync_at)],
                ],
            )
        )
        sheets.append(
            Sheet(
                "By district",
                ["District", "Total", "Resolved", "Unresolved", "Overdue"],
                [[b.label, b.total, b.resolved, b.unresolved, b.overdue] for b in ds.by_district(db, user, f)],
            )
        )
        sheets.append(
            Sheet(
                "Trend",
                ["Day", "Received", "Resolved"],
                [[t.period, t.received, t.resolved] for t in ds.trend(db, user, f)],
            )
        )

    elif kind in ("district_register", "full_export"):
        rows, _ = ds.error_list(db, user, f, page=1, page_size=100000)
        sheets.append(
            Sheet(
                "Errors",
                [
                    "Error ID", "District", "Team", "EA", "Category", "Supervisor", "Enumerator", "Description",
                    "Date received", "Support", "Status", "Next follow-up", "Overdue", "Field Monitor",
                    "Last action", "Follow-ups",
                ],
                [
                    [
                        r.display_id, r.district, r.team or "", r.ea or "", r.category, r.supervisor_name or "",
                        r.enumerator_name or "", r.description, r.date_received.isoformat(), r.support_method.value,
                        r.status.value, _fmt(r.next_follow_up_at), "Yes" if r.overdue else "", r.monitor,
                        _fmt(r.last_action_at), r.follow_up_count,
                    ]
                    for r in rows
                ],
                landscape=True,
            )
        )
        if kind == "full_export":
            q = ds.base_query(db, user, f).options(selectinload(ErrorRecord.follow_ups), selectinload(ErrorRecord.activity))
            records = db.execute(q).scalars().all()
            users = {u.id: u.full_name for u in db.execute(select(User)).scalars()}
            fu_rows, act_rows = [], []
            for r in records:
                for fu in r.follow_ups:
                    fu_rows.append([r.display_id, _fmt(fu.at), fu.method.value, fu.contacted or "", fu.outcome or "", fu.comments or "", fu.lat or "", fu.lng or "", users.get(fu.user_id, "")])
                for a in r.activity:
                    act_rows.append([r.display_id, _fmt(a.client_at), a.previous_status.value if a.previous_status else "", a.new_status.value, a.action_taken or "", a.comments or "", users.get(a.user_id, "")])
            sheets.append(Sheet("Follow-ups", ["Error ID", "At", "Method", "Contacted", "Outcome", "Comments", "Lat", "Lng", "By"], fu_rows, landscape=True))
            sheets.append(Sheet("Activity", ["Error ID", "At", "Previous", "New", "Action taken", "Comments", "By"], act_rows, landscape=True))

    elif kind == "overdue":
        rows = ds.overdue_list(db, user, f)
        sheets.append(
            Sheet(
                "Overdue",
                ["Error ID", "District", "Team", "Category", "Supervisor", "Field Monitor", "Follow-up due", "Hours overdue"],
                [[r.display_id, r.district, r.team or "", r.category, r.supervisor_name or "", r.monitor, _fmt(r.next_follow_up_at), r.hours_overdue] for r in rows],
                landscape=True,
            )
        )

    elif kind == "category_analysis":
        buckets = ds.by_category(db, user, f)
        sheets.append(
            Sheet(
                "By category",
                ["Category", "Total", "Resolved", "Unresolved", "Overdue", "Resolution rate (%)"],
                [[b.label, b.total, b.resolved, b.unresolved, b.overdue, round(b.resolved / b.total * 100, 1) if b.total else 0] for b in buckets],
            )
        )
        # Category x district matrix
        rows = db.execute(ds.base_query(db, user, f)).scalars().all()
        cats = {c.id: c.name for c in db.execute(select(ErrorCategory)).scalars()}
        dists = {d.id: d.name for d in db.execute(select(District)).scalars()}
        matrix: dict[tuple, int] = {}
        for r in rows:
            matrix[(r.category_id, r.district_id)] = matrix.get((r.category_id, r.district_id), 0) + 1
        d_ids = sorted({d for _, d in matrix}, key=lambda i: dists.get(i, ""))
        c_ids = sorted({c for c, _ in matrix}, key=lambda i: cats.get(i, ""))
        sheets.append(
            Sheet(
                "Category by district",
                ["Category"] + [dists.get(d, str(d)) for d in d_ids],
                [[cats.get(c, str(c))] + [matrix.get((c, d), 0) for d in d_ids] for c in c_ids],
                landscape=True,
            )
        )

    elif kind == "team_performance":
        rows = ds.by_team(db, user, f)
        sheets.append(
            Sheet(
                "Teams",
                ["Team", "District", "Supervisor", "Total", "Resolved", "Unresolved", "Overdue"],
                [[r.team, r.district, r.supervisor or "", r.total, r.resolved, r.unresolved, r.overdue] for r in rows],
                landscape=True,
            )
        )

    elif kind == "monitor_activity":
        rows = ds.by_monitor(db, user, f)
        sheets.append(
            Sheet(
                "Field Monitors",
                ["Field Monitor", "Districts", "Total", "Unresolved", "Overdue", "Last sync", "Last activity", "App version", "Pending on tablet"],
                [[r.full_name, r.districts, r.total, r.unresolved, r.overdue, _fmt(r.last_sync_at), _fmt(r.last_activity_at), r.app_version or "", r.pending_reported] for r in rows],
                landscape=True,
            )
        )

    elif kind == "error_detail":
        if not error_id:
            raise ValueError("error_id is required for the error detail sheet")
        q = ds.base_query(db, user, f).where(ErrorRecord.id == error_id).options(selectinload(ErrorRecord.follow_ups), selectinload(ErrorRecord.activity))
        r = db.execute(q).scalars().first()
        if r is None:
            raise ValueError("Error not found in your scope")
        users = {u.id: u.full_name for u in db.execute(select(User)).scalars()}
        dist = db.get(District, r.district_id)
        team = db.get(Team, r.team_id) if r.team_id else None
        cat = db.get(ErrorCategory, r.category_id)
        sheets.append(
            Sheet(
                "Error",
                ["Field", "Value"],
                [
                    ["Error ID", r.display_id], ["District", dist.name if dist else ""], ["Team", f"{team.code} {team.name}" if team else ""],
                    ["Supervisor", r.supervisor_name or ""], ["Enumerator", r.enumerator_name or ""], ["Category", cat.name if cat else ""],
                    ["Description", r.description], ["Date received", r.date_received.isoformat()], ["Support method", r.support_method.value],
                    ["Action taken", r.action_taken or ""], ["Comments", r.comments or ""], ["Status", r.status.value],
                    ["Resolved at", _fmt(r.resolved_at)], ["Next follow-up", _fmt(r.next_follow_up_at)], ["GPS", f"{r.lat}, {r.lng} (±{r.accuracy_m} m)" if r.lat else ""],
                    ["Field Monitor", users.get(r.user_id, "")],
                ],
            )
        )
        sheets.append(Sheet("Follow-ups", ["At", "Method", "Contacted", "Outcome", "Comments"], [[_fmt(fu.at), fu.method.value, fu.contacted or "", fu.outcome or "", fu.comments or ""] for fu in sorted(r.follow_ups, key=lambda x: x.at)]))
        sheets.append(Sheet("Activity", ["At", "Previous", "New", "Action taken", "Comments"], [[_fmt(a.client_at), a.previous_status.value if a.previous_status else "", a.new_status.value, a.action_taken or "", a.comments or ""] for a in sorted(r.activity, key=lambda x: x.client_at)]))

    return Report(kind=kind, title=REPORT_KINDS[kind], sheets=sheets, generated_at=now, filters=_filters_text(f))


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(bold=True, color="FFFFFF")
    for sheet in report.sheets:
        ws = wb.create_sheet(sheet.name[:31])
        ws.append([report.title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Generated {report.generated_at.strftime('%Y-%m-%d %H:%M UTC')} · Filters: {report.filters}"])
        ws.append([])
        ws.append(sheet.headers)
        for cell in ws[4]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.rows:
            ws.append([_fmt(v) for v in row])
        for i, h in enumerate(sheet.headers, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in sheet.rows[:200] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report: Report) -> bytes:
    buf = io.BytesIO()
    is_landscape = any(s.landscape for s in report.sheets)
    pagesize = landscape(A4) if is_landscape else A4
    doc = SimpleDocTemplate(
        buf, pagesize=pagesize, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title=report.title,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"SLPHC 2026 · {report.title}", styles["Title"]),
        Paragraph(f"Generated {report.generated_at.strftime('%Y-%m-%d %H:%M UTC')} · Filters: {report.filters}", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    cell_style = styles["BodyText"]
    cell_style.fontSize = 8
    cell_style.leading = 10
    for sheet in report.sheets:
        story.append(Paragraph(sheet.name, styles["Heading2"]))
        data = [[Paragraph(str(h), cell_style) for h in sheet.headers]]
        for row in sheet.rows:
            data.append([Paragraph(str(_fmt(v)), cell_style) for v in row])
        if len(data) == 1:
            data.append([Paragraph("No records", cell_style)] + [""] * (len(sheet.headers) - 1))
        avail = pagesize[0] - 24 * mm
        col_w = avail / max(len(sheet.headers), 1)
        t = Table(data, colWidths=[col_w] * len(sheet.headers), repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(t)
        story.append(Spacer(1, 6 * mm))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(pagesize[0] - 12 * mm, 8 * mm, f"Page {doc_.page}")
        canvas.drawString(12 * mm, 8 * mm, "SLPHC 2026 Field Monitor Error Follow-up")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
