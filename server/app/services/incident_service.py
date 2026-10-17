"""Census Incident Log: reporting (web and public SA-code page), handling actions with a dated log,
the register with scoping, the dashboard figures and the exports."""

from __future__ import annotations

import csv
import hashlib
import io
import secrets
import statistics
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.core import incidents as spec
from app.core.districts import district_abbreviation
from app.models import District, EnumerationArea, Incident, IncidentSaCode, IncidentUpdate, Region, Team, User, UserScope
from app.models.user import Role
from app.schemas.incident import (
    ActionIn,
    CategoryRow,
    CountRow,
    Dashboard,
    DistrictRow,
    IncidentEdit,
    IncidentIn,
    IncidentOut,
    OwnerRow,
    PersonOut,
    PublicIncidentIn,
    PublicIncidentOut,
    PublicTrackOut,
    SaCodeImportOut,
    SaCodeIn,
    SaCodeOut,
    SaRow,
    TrendRow,
    UpdateOut,
)
from app.services import report_service
from app.services.reference import current_settings
from app.services.scope import district_ids_for


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _aware(dt: datetime | None) -> datetime | None:
    return None if dt is None else (dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc))


def _hash(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def _hours(a: datetime | None, b: datetime | None) -> float | None:
    a, b = _aware(a), _aware(b)
    return round((b - a).total_seconds() / 3600, 1) if a and b else None


def _pct(n: int, d: int) -> float:
    return round(100.0 * n / d, 1) if d else 0.0


# ---- numbering, defaults ---------------------------------------------------------------------------


def _number(db: Session, district: District | None, when: datetime) -> str:
    abbr = district_abbreviation(district.code) if district else "HQ"
    prefix = f"INC-{abbr}-{when:%Y%m}-"
    last = db.execute(select(Incident.number).where(Incident.number.like(prefix + "%")).order_by(Incident.number.desc()).limit(1)).scalar()
    n = int(last.rsplit("-", 1)[1]) + 1 if last else 1
    return f"{prefix}{n:04d}"


def due_hours(db: Session) -> dict[str, int]:
    s = current_settings(db)
    return {sev: int(s.get(f"incident_due_hours_{sev.lower()}") or spec.DEFAULT_DUE_HOURS[sev]) for sev in spec.DEFAULT_DUE_HOURS}


def default_owner(db: Session, district_id: int | None) -> User | None:
    """The District Statistician of the district (the first active one), else nobody."""
    if district_id is None:
        return None
    return db.execute(
        select(User).join(UserScope, UserScope.user_id == User.id)
        .where(User.role == Role.DISTRICT_STATISTICIAN, User.active.is_(True), UserScope.district_id == district_id).order_by(User.id)
    ).scalars().first()


def _resolve_sa(db: Session, sa_code: str | None) -> tuple[IncidentSaCode | None, Team | None, District | None, bool]:
    """The uploaded SA code, else a team whose code matches the digits; recognised=False when neither."""
    if not sa_code:
        return None, None, None, True
    row = db.execute(select(IncidentSaCode).where(IncidentSaCode.sa_code == sa_code, IncidentSaCode.active.is_(True))).scalars().first()
    if row is not None:
        team = db.get(Team, row.team_id) if row.team_id else None
        district = db.get(District, row.district_id) if row.district_id else (db.get(District, team.district_id) if team else None)
        return row, team, district, True
    digits = sa_code.lstrip("0") or "0"
    team = db.execute(select(Team).where(Team.active.is_(True), func.ltrim(Team.code, "0") == digits)).scalars().first()
    if team is None and len(digits) >= 4:
        team = db.execute(select(Team).where(Team.active.is_(True), Team.code.like(f"%{digits}"))).scalars().first()
    if team is not None:
        return None, team, db.get(District, team.district_id), True
    return None, None, None, False


def _log(db: Session, inc: Incident, type_: str, text: str | None = None, old: str | None = None, new: str | None = None, user: User | None = None, author_name: str | None = None) -> None:
    db.add(IncidentUpdate(incident_id=inc.id, at=_now(), author_id=user.id if user else None, author_name=user.full_name if user else author_name, type=type_, text=text, old_value=old, new_value=new))


# ---- reporting ---------------------------------------------------------------------------------------


def _create(db: Session, data: dict, *, user: User | None, district: District | None, team: Team | None, sa_code: str | None, recognised: bool, reporter: dict) -> Incident:
    errors = spec.validate(data)
    if errors:
        raise HTTPException(422, {"message": "Please check the report", "errors": errors})
    now = _now()
    occurred = _aware(data["occurred_at"])
    if occurred > now + timedelta(minutes=5):
        raise HTTPException(400, "The incident time cannot be in the future")
    owner = default_owner(db, district.id if district else None)
    inc = Incident(
        number=_number(db, district, now), reported_by=user.id if user else None, reporter_name=reporter.get("name"), reporter_role=reporter.get("role"), reporter_phone=reporter.get("phone"), reporter_code=reporter.get("code"),
        sa_code=sa_code, sa_recognised=recognised, district_id=district.id if district else None, team_id=team.id if team else None,
        ea_code=(data.get("ea_code") or "").strip() or None, locality=(data.get("locality") or "").strip() or None, lat=data.get("lat"), lng=data.get("lng"),
        phase=data["phase"], occurred_at=occurred, reported_at=now, category=data["category"], subcategory=data.get("subcategory") or None, severity=data["severity"],
        title=data["title"].strip(), description=data["description"].strip(), action_taken=(data.get("action_taken") or "").strip() or None, help_needed=(data.get("help_needed") or "").strip() or None,
        teams_affected=data.get("teams_affected"), staff_affected=data.get("staff_affected"), households_affected=data.get("households_affected"), days_lost=data.get("days_lost"),
        asset_tags=(data.get("asset_tags") or "").strip() or None, reference=(data.get("reference") or "").strip() or None,
        status="NEW", level="DISTRICT" if district else "NATIONAL", owner_id=owner.id if owner else None, due_at=now + timedelta(hours=due_hours(db)[data["severity"]]),
    )
    db.add(inc)
    db.flush()
    _log(db, inc, "CREATED", text=f"Reported{' by ' + reporter['name'] if reporter.get('name') else ''}", user=user, author_name=reporter.get("name"))
    if owner:
        _log(db, inc, "ASSIGN", text="Assigned to the District Statistician by default", new=owner.full_name)
    return inc


def report(db: Session, user: User, body: IncidentIn) -> Incident:
    scope = district_ids_for(db, user)
    district = db.get(District, body.district_id) if body.district_id else None
    team = db.get(Team, body.team_id) if body.team_id else None
    if team and district is None:
        district = db.get(District, team.district_id)
    sa_row, sa_team, sa_district, recognised = _resolve_sa(db, (body.sa_code or "").strip() or None)
    if team is None and sa_team is not None:
        team = sa_team
    if district is None and sa_district is not None:
        district = sa_district
    if district is None and scope is not None and len(scope) == 1:
        district = db.get(District, scope[0])
    if district is None:
        raise HTTPException(400, "Choose the district")
    if scope is not None and district.id not in scope:
        raise HTTPException(403, "You can only report incidents in your own district")
    reporter = {"name": (body.reporter_name or "").strip() or None, "role": body.reporter_role, "phone": (body.reporter_phone or "").strip() or None, "code": (body.reporter_code or "").strip() or None}
    return _create(db, body.model_dump(), user=user, district=district, team=team, sa_code=(body.sa_code or "").strip() or None, recognised=recognised, reporter=reporter)


def public_report(db: Session, body: PublicIncidentIn) -> tuple[Incident, str]:
    role = body.reporter_role if body.reporter_role in dict(spec.REPORTER_ROLES) else "OTHER"
    required = spec.ROLE_REQUIRED.get(role, [])
    sa_code = (body.sa_code or "").strip() or None
    code = (body.reporter_code or "").strip() or None
    errors = []
    if "code" in required and not code:
        errors.append("Enter your application code")
    if "sa" in required and not sa_code:
        errors.append("Enter your 12-digit SA code")
    if sa_code and not spec.SA_CODE_RE.match(sa_code):
        errors.append("The SA code has 12 digits")
    if errors:
        raise HTTPException(422, {"message": "Please check the report", "errors": errors})
    sa_row, team, district, recognised = svc_resolve(db, sa_code)
    if district is None and body.district:
        district = _district_by(db, body.district)
        recognised = True if district is not None and sa_code is None else recognised
    if "district" in required and district is None and sa_code is None:
        raise HTTPException(422, {"message": "Please check the report", "errors": ["Choose your district"]})
    token = secrets.token_urlsafe(18)
    reporter = {"name": body.reporter_name.strip(), "role": role, "phone": body.reporter_phone.strip(), "code": code}
    inc = _create(db, body.model_dump(), user=None, district=district, team=team, sa_code=sa_code, recognised=recognised, reporter=reporter)
    inc.track_token_hash = _hash(token)
    if not recognised:
        _log(db, inc, "COMMENT", text=f"SA code {sa_code} is not in the uploaded list: the district must be set by a manager")
    db.flush()
    return inc, token


def svc_resolve(db: Session, sa_code: str | None):
    return _resolve_sa(db, sa_code)


def public_track(db: Session, number: str, token: str) -> PublicTrackOut:
    inc = db.execute(select(Incident).options(selectinload(Incident.updates)).where(Incident.number == number.strip().upper(), Incident.deleted.is_(False))).scalars().first()
    if inc is None or not inc.track_token_hash or inc.track_token_hash != _hash(token):
        raise HTTPException(404, "No incident with this number and tracking code")
    district = db.get(District, inc.district_id) if inc.district_id else None
    return PublicTrackOut(
        number=inc.number, title=inc.title, status=inc.status, status_label=spec.label("status", inc.status), severity=inc.severity, district=district.name if district else None,
        reported_at=inc.reported_at, acknowledged_at=inc.acknowledged_at, resolved_at=inc.resolved_at, resolution=inc.resolution,
        updates=[_update_out(u) for u in inc.updates if u.type in ("STATUS", "REOPEN", "CREATED")],
    )


# ---- handling ----------------------------------------------------------------------------------------


def get(db: Session, user: User, incident_id: int) -> Incident:
    inc = db.execute(select(Incident).options(selectinload(Incident.updates)).where(Incident.id == incident_id, Incident.deleted.is_(False))).scalars().first()
    if inc is None:
        raise HTTPException(404, "Incident not found")
    scope = district_ids_for(db, user)
    if scope is not None and inc.district_id not in scope:
        raise HTTPException(403, "This incident is outside your district")
    return inc


def edit(db: Session, user: User, inc: Incident, body: IncidentEdit) -> Incident:
    data = body.model_dump(exclude_unset=True)
    check = {"category": data.get("category", inc.category), "subcategory": data.get("subcategory", inc.subcategory), "severity": data.get("severity", inc.severity), "phase": data.get("phase", inc.phase), "description": data.get("description", inc.description)}
    errors = spec.validate(check)
    if errors:
        raise HTTPException(422, {"message": "Please check the changes", "errors": errors})
    if "district_id" in data and data["district_id"]:
        scope = district_ids_for(db, user)
        if scope is not None and data["district_id"] not in scope:
            raise HTTPException(403, "You cannot move an incident outside your district")
        if db.get(District, data["district_id"]) is None:
            raise HTTPException(404, "District not found")
        if inc.district_id != data["district_id"]:
            old = db.get(District, inc.district_id).name if inc.district_id else "none"
            _log(db, inc, "COMMENT", text="District changed", old=old, new=db.get(District, data["district_id"]).name, user=user)
            if inc.owner_id is None:
                owner = default_owner(db, data["district_id"])
                if owner:
                    inc.owner_id = owner.id
                    _log(db, inc, "ASSIGN", text="Assigned to the District Statistician", new=owner.full_name, user=user)
            inc.sa_recognised = True
    if "severity" in data and data["severity"] != inc.severity:
        _log(db, inc, "SEVERITY", old=inc.severity, new=data["severity"], user=user)
    for k, v in data.items():
        setattr(inc, k, v.strip() if isinstance(v, str) else v)
    db.flush()
    return inc


def act(db: Session, user: User, inc: Incident, body: ActionIn) -> Incident:
    now = _now()
    a = body.action
    text = (body.text or "").strip() or None
    if a == "comment":
        if not text:
            raise HTTPException(400, "Write the comment")
        _log(db, inc, "COMMENT", text=text, user=user)
    elif a == "acknowledge":
        if inc.status != "NEW":
            raise HTTPException(409, "Only a new incident can be acknowledged")
        inc.status, inc.acknowledged_at = "ACKNOWLEDGED", now
        if inc.owner_id is None:
            inc.owner_id = user.id
        _log(db, inc, "STATUS", text=text, old="NEW", new="ACKNOWLEDGED", user=user)
    elif a == "assign":
        owner = db.get(User, body.owner_id) if body.owner_id else None
        if owner is None or not owner.active:
            raise HTTPException(404, "Owner not found")
        old = db.get(User, inc.owner_id).full_name if inc.owner_id else None
        inc.owner_id = owner.id
        if body.due_at:
            inc.due_at = _aware(body.due_at)
        if inc.status == "NEW":
            inc.status, inc.acknowledged_at = "ACKNOWLEDGED", now
        _log(db, inc, "ASSIGN", text=text, old=old, new=owner.full_name, user=user)
    elif a == "start":
        if inc.status not in ("NEW", "ACKNOWLEDGED"):
            raise HTTPException(409, "The incident is not waiting to start")
        if inc.acknowledged_at is None:
            inc.acknowledged_at = now
        _log(db, inc, "STATUS", text=text, old=inc.status, new="IN_PROGRESS", user=user)
        inc.status = "IN_PROGRESS"
    elif a == "resolve":
        if inc.status not in spec.OPEN_STATUSES:
            raise HTTPException(409, "The incident is not open")
        if not text:
            raise HTTPException(400, "Say what was done to resolve it")
        inc.resolution, inc.resolved_at, inc.root_cause_fixed = text, now, body.root_cause_fixed
        if body.lessons:
            inc.lessons = body.lessons.strip()
        if inc.acknowledged_at is None:
            inc.acknowledged_at = now
        _log(db, inc, "STATUS", text=text, old=inc.status, new="RESOLVED", user=user)
        inc.status = "RESOLVED"
    elif a == "close":
        if inc.status != "RESOLVED":
            raise HTTPException(409, "Resolve the incident before closing it")
        inc.status, inc.closed_at = "CLOSED", now
        _log(db, inc, "STATUS", text=text, old="RESOLVED", new="CLOSED", user=user)
    elif a == "reject":
        if inc.status not in spec.OPEN_STATUSES:
            raise HTTPException(409, "Only an open incident can be rejected")
        if not text:
            raise HTTPException(400, "Give the reason for rejecting it")
        inc.reject_reason, inc.closed_at = text, now
        _log(db, inc, "STATUS", text=text, old=inc.status, new="REJECTED", user=user)
        inc.status = "REJECTED"
    elif a == "reopen":
        if inc.status not in ("RESOLVED", "CLOSED", "REJECTED"):
            raise HTTPException(409, "The incident is already open")
        if not text:
            raise HTTPException(400, "Say why it is reopened")
        _log(db, inc, "REOPEN", text=text, old=inc.status, new="IN_PROGRESS", user=user)
        inc.status, inc.resolved_at, inc.closed_at, inc.reject_reason = "IN_PROGRESS", None, None, None
    elif a == "escalate":
        level = body.level or {"DISTRICT": "REGION", "REGION": "NATIONAL"}.get(inc.level, "NATIONAL")
        if level == inc.level:
            raise HTTPException(409, f"The incident is already at {spec.label('level', level).lower()} level")
        _log(db, inc, "ESCALATE", text=text, old=inc.level, new=level, user=user)
        inc.level = level
        if inc.status == "NEW":
            inc.status, inc.acknowledged_at = "ACKNOWLEDGED", now
    db.flush()
    return inc


def delete(db: Session, inc: Incident) -> None:
    inc.deleted = True
    db.flush()


# ---- reads -------------------------------------------------------------------------------------------


def _update_out(u: IncidentUpdate) -> UpdateOut:
    return UpdateOut(id=u.id, at=u.at, author=u.author_name, type=u.type, type_label=spec.label("update", u.type), text=u.text, old_value=u.old_value, new_value=u.new_value)


class Names:
    def __init__(self, db: Session, rows: list[Incident]):
        d_ids = {x.district_id for x in rows if x.district_id}
        self.districts = {d.id: d for d in db.execute(select(District).where(District.id.in_(d_ids))).scalars()} if d_ids else {}
        self.regions = {r.id: r.name for r in db.execute(select(Region)).scalars()}
        t_ids = {x.team_id for x in rows if x.team_id}
        self.teams = {t.id: t.code for t in db.execute(select(Team).where(Team.id.in_(t_ids))).scalars()} if t_ids else {}
        u_ids = {x.owner_id for x in rows if x.owner_id} | {x.reported_by for x in rows if x.reported_by}
        self.users = {u.id: u.full_name for u in db.execute(select(User).where(User.id.in_(u_ids))).scalars()} if u_ids else {}


def to_out(inc: Incident, names: Names, with_updates: bool = False) -> IncidentOut:
    now = _now()
    district = names.districts.get(inc.district_id)
    open_ = inc.status in spec.OPEN_STATUSES
    due = _aware(inc.due_at)
    return IncidentOut(
        id=inc.id, number=inc.number, status=inc.status, status_label=spec.label("status", inc.status), severity=inc.severity, severity_label=spec.label("severity", inc.severity),
        category=inc.category, category_label=spec.label("category", inc.category), subcategory=inc.subcategory, subcategory_label=spec.SUBCATEGORY_LABEL.get((inc.category, inc.subcategory or ""), inc.subcategory),
        phase=inc.phase, phase_label=spec.label("phase", inc.phase), title=inc.title, description=inc.description, action_taken=inc.action_taken, help_needed=inc.help_needed,
        district_id=inc.district_id, district=district.name if district else None, region=names.regions.get(district.region_id) if district else None,
        team_id=inc.team_id, team_code=names.teams.get(inc.team_id), sa_code=inc.sa_code, sa_recognised=inc.sa_recognised, ea_code=inc.ea_code, locality=inc.locality, lat=inc.lat, lng=inc.lng,
        occurred_at=inc.occurred_at, reported_at=inc.reported_at, reported_by_name=names.users.get(inc.reported_by), reporter_name=inc.reporter_name, reporter_role=inc.reporter_role,
        reporter_role_label=spec.label("reporter_role", inc.reporter_role) if inc.reporter_role else None, reporter_phone=inc.reporter_phone, reporter_code=inc.reporter_code,
        teams_affected=inc.teams_affected, staff_affected=inc.staff_affected, households_affected=inc.households_affected, days_lost=inc.days_lost, asset_tags=inc.asset_tags, reference=inc.reference,
        level=inc.level, owner_id=inc.owner_id, owner_name=names.users.get(inc.owner_id), due_at=inc.due_at, overdue=bool(open_ and due and due < now),
        unacknowledged_hours=_hours(inc.reported_at, now) if inc.status == "NEW" else None,
        acknowledged_at=inc.acknowledged_at, resolved_at=inc.resolved_at, closed_at=inc.closed_at, resolution=inc.resolution, root_cause_fixed=inc.root_cause_fixed, lessons=inc.lessons, reject_reason=inc.reject_reason,
        hours_to_acknowledge=_hours(inc.reported_at, inc.acknowledged_at), hours_to_resolve=_hours(inc.reported_at, inc.resolved_at),
        updates=[_update_out(u) for u in inc.updates] if with_updates else [],
    )


def _query(db: Session, user: User, f: dict):
    q = select(Incident).where(Incident.deleted.is_(False))
    scope = district_ids_for(db, user)
    if scope is not None:
        q = q.where(Incident.district_id.in_(scope))
    if f.get("district_id"):
        q = q.where(Incident.district_id == f["district_id"])
    if f.get("region_id"):
        q = q.where(Incident.district_id.in_(select(District.id).where(District.region_id == f["region_id"])))
    if f.get("status"):
        q = q.where(Incident.status == f["status"]) if f["status"] != "OPEN" else q.where(Incident.status.in_(spec.OPEN_STATUSES))
    for k in ("severity", "category", "phase", "level", "owner_id", "team_id"):
        if f.get(k):
            q = q.where(getattr(Incident, k) == f[k])
    if f.get("date_from"):
        q = q.where(Incident.reported_at >= datetime.combine(f["date_from"], datetime.min.time(), tzinfo=timezone.utc))
    if f.get("date_to"):
        q = q.where(Incident.reported_at < datetime.combine(f["date_to"] + timedelta(days=1), datetime.min.time(), tzinfo=timezone.utc))
    if f.get("search"):
        s = f"%{f['search'].strip().lower()}%"
        q = q.where(func.lower(Incident.number).like(s) | func.lower(Incident.title).like(s) | func.lower(Incident.description).like(s) | func.lower(func.coalesce(Incident.reporter_name, "")).like(s) | func.lower(func.coalesce(Incident.sa_code, "")).like(s))
    return q


def list_incidents(db: Session, user: User, f: dict, limit: int = 500) -> list[IncidentOut]:
    rows = db.execute(_query(db, user, f).order_by(Incident.reported_at.desc()).limit(limit)).scalars().all()
    names = Names(db, rows)
    out = [to_out(x, names) for x in rows]
    if f.get("overdue"):
        out = [x for x in out if x.overdue]
    return out


def people(db: Session, user: User, district_id: int | None) -> list[PersonOut]:
    """Accounts an incident can be assigned to: District Statisticians and managers in scope."""
    rows = db.execute(select(User).where(User.active.is_(True), User.role.in_([Role.DISTRICT_STATISTICIAN, Role.DISTRICT_DQM, Role.REGIONAL, Role.NATIONAL_DQM, Role.ADMIN, Role.ME])).order_by(User.full_name)).scalars().all()
    out = []
    for u in rows:
        ids = district_ids_for(db, u)
        if district_id is not None and ids is not None and district_id not in ids:
            continue
        out.append(PersonOut(id=u.id, full_name=u.full_name, role=u.role.value, district_ids=ids or []))
    return out


def open_for_team(db: Session, user: User, team_id: int) -> list[IncidentOut]:
    rows = db.execute(_query(db, user, {"team_id": team_id, "status": "OPEN"}).order_by(Incident.severity, Incident.reported_at)).scalars().all()
    return [to_out(x, Names(db, rows)) for x in rows]


# ---- dashboard ---------------------------------------------------------------------------------------


def dashboard(db: Session, user: User, f: dict) -> Dashboard:
    rows = db.execute(_query(db, user, f).options(selectinload(Incident.updates)).order_by(Incident.reported_at.desc())).scalars().all()
    names = Names(db, rows)
    outs = [to_out(x, names) for x in rows]
    now = _now()
    today = now.date()
    open_ = [x for x in outs if x.status in spec.OPEN_STATUSES]
    n = len(outs)

    def counts(key, pairs):
        c = defaultdict(int)
        for x in outs:
            c[getattr(x, key)] += 1
        return [CountRow(code=code, label=lb, count=c[code], pct=_pct(c[code], n)) for code, lb in pairs]

    # by district: every district in scope, even with no incident
    scope = district_ids_for(db, user)
    dq = select(District).order_by(District.name)
    if scope is not None:
        dq = dq.where(District.id.in_(scope))
    districts = db.execute(dq).scalars().all()
    ea_counts = dict(db.execute(select(Team.district_id, func.count()).select_from(EnumerationArea).join(Team, Team.id == EnumerationArea.team_id).where(EnumerationArea.active.is_(True)).group_by(Team.district_id)).all())
    by_district = []
    for d in districts:
        xs = [x for x in outs if x.district_id == d.id]
        res = [x.hours_to_resolve for x in xs if x.hours_to_resolve is not None]
        by_district.append(DistrictRow(
            district_id=d.id, district=d.name, region=names.regions.get(d.region_id, ""), open=sum(1 for x in xs if x.status in spec.OPEN_STATUSES), new=sum(1 for x in xs if x.status == "NEW"),
            critical_open=sum(1 for x in xs if x.severity == "CRITICAL" and x.status in spec.OPEN_STATUSES), overdue=sum(1 for x in xs if x.overdue), resolved=sum(1 for x in xs if x.status in ("RESOLVED", "CLOSED")),
            total=len(xs), per_100_eas=round(100 * len(xs) / ea_counts[d.id], 2) if ea_counts.get(d.id) else None, median_hours_to_resolve=round(statistics.median(res), 1) if res else None,
        ))
    # categories with 7-day movement
    by_category = []
    for code, lb, _subs, _sev in spec.CATEGORIES:
        xs = [x for x in outs if x.category == code]
        by_category.append(CategoryRow(
            category=code, label=lb, total=len(xs), open=sum(1 for x in xs if x.status in spec.OPEN_STATUSES),
            critical=sum(1 for x in xs if x.severity == "CRITICAL"), major=sum(1 for x in xs if x.severity == "MAJOR"), minor=sum(1 for x in xs if x.severity == "MINOR"),
            last_7_days=sum(1 for x in xs if _aware(x.reported_at) >= now - timedelta(days=7)), previous_7_days=sum(1 for x in xs if now - timedelta(days=14) <= _aware(x.reported_at) < now - timedelta(days=7)),
        ))
    sub = defaultdict(int)
    for x in outs:
        if x.subcategory:
            sub[(x.category, x.subcategory)] += 1
    by_sub = sorted([CountRow(code=f"{c}:{s}", label=f"{spec.CATEGORY_LABEL[c]} · {spec.SUBCATEGORY_LABEL.get((c, s), s)}", count=v, pct=_pct(v, n)) for (c, s), v in sub.items()], key=lambda r: -r.count)[:12]
    # trend over the last 30 days
    start = today - timedelta(days=29)
    reported = defaultdict(int)
    resolved = defaultdict(int)
    for x in outs:
        reported[_aware(x.reported_at).date()] += 1
        if x.resolved_at:
            resolved[_aware(x.resolved_at).date()] += 1
    trend = []
    open_before = sum(1 for x in outs if _aware(x.reported_at).date() < start) - sum(1 for x in outs if x.resolved_at and _aware(x.resolved_at).date() < start) - sum(1 for x in outs if x.status == "REJECTED" and x.closed_at and _aware(x.closed_at).date() < start)
    running = open_before
    for i in range(30):
        day = start + timedelta(days=i)
        running += reported[day] - resolved[day]
        trend.append(TrendRow(day=day, reported=reported[day], resolved=resolved[day], open_end_of_day=max(running, 0)))
    # owners, SAs
    owners = defaultdict(lambda: {"open": 0, "overdue": 0, "resolved": 0, "name": ""})
    for x in outs:
        o = owners[x.owner_id]
        o["name"] = x.owner_name or "Unassigned"
        if x.status in spec.OPEN_STATUSES:
            o["open"] += 1
        if x.overdue:
            o["overdue"] += 1
        if x.status in ("RESOLVED", "CLOSED"):
            o["resolved"] += 1
    by_owner = sorted([OwnerRow(owner_id=k, owner=v["name"], open=v["open"], overdue=v["overdue"], resolved=v["resolved"]) for k, v in owners.items()], key=lambda r: (-r.open, r.owner))
    sas = defaultdict(lambda: {"n": 0, "open": 0, "critical": 0, "team": None, "district": None})
    for x in outs:
        key = x.sa_code or x.team_code
        if not key:
            continue
        s = sas[key]
        s["n"] += 1
        s["open"] += x.status in spec.OPEN_STATUSES
        s["critical"] += x.severity == "CRITICAL"
        s["team"], s["district"] = x.team_code, x.district
    top_sas = sorted([SaRow(sa_code=k, team_code=v["team"], district=v["district"], incidents=v["n"], open=v["open"], critical=v["critical"]) for k, v in sas.items()], key=lambda r: (-r.open, -r.incidents))[:10]
    ack = [x.hours_to_acknowledge for x in outs if x.hours_to_acknowledge is not None]
    res = [x.hours_to_resolve for x in outs if x.hours_to_resolve is not None]
    sev_rank = {"CRITICAL": 0, "MAJOR": 1, "MINOR": 2}
    needs = sorted([x for x in open_ if x.status == "NEW" or x.severity == "CRITICAL" or x.overdue], key=lambda x: (sev_rank[x.severity], x.reported_at))[:25]
    if scope is None:
        scope_label = "national"
    elif len({d.region_id for d in districts}) == 1 and len(districts) > 1:
        scope_label = f"region: {names.regions.get(districts[0].region_id, '')}"
    else:
        scope_label = ", ".join(d.name for d in districts) or "no district assigned"
    return Dashboard(
        scope=scope_label, filters={k: (str(v) if v is not None else None) for k, v in f.items()},
        total=n, open=len(open_), new_unacknowledged=sum(1 for x in outs if x.status == "NEW"), critical_open=sum(1 for x in open_ if x.severity == "CRITICAL"),
        overdue=sum(1 for x in outs if x.overdue), due_today=sum(1 for x in open_ if x.due_at and _aware(x.due_at).date() == today and not x.overdue),
        reported_today=sum(1 for x in outs if _aware(x.reported_at).date() == today), reported_7_days=sum(1 for x in outs if _aware(x.reported_at) >= now - timedelta(days=7)),
        resolved_7_days=sum(1 for x in outs if x.resolved_at and _aware(x.resolved_at) >= now - timedelta(days=7)),
        closed=sum(1 for x in outs if x.status == "CLOSED"), rejected=sum(1 for x in outs if x.status == "REJECTED"), unrecognised_sa=sum(1 for x in open_ if not x.sa_recognised),
        hours_to_acknowledge_median=round(statistics.median(ack), 1) if ack else None, hours_to_resolve_median=round(statistics.median(res), 1) if res else None,
        by_status=counts("status", spec.STATUSES), by_severity=counts("severity", spec.SEVERITIES), by_phase=counts("phase", spec.PHASES), by_level=counts("level", spec.LEVELS),
        by_reporter_role=counts("reporter_role", spec.REPORTER_ROLES), by_category=by_category, by_subcategory=by_sub, by_district=by_district, by_owner=by_owner, top_sas=top_sas,
        trend=trend, needs_action=needs, recent=outs[:15],
        points=[{"lat": x.lat, "lng": x.lng, "severity": x.severity, "number": x.number, "status": x.status, "title": x.title} for x in outs if x.lat is not None and x.lng is not None][:500],
    )


# ---- SA codes ----------------------------------------------------------------------------------------


def sa_codes(db: Session, user: User) -> list[SaCodeOut]:
    scope = district_ids_for(db, user)
    q = select(IncidentSaCode).order_by(IncidentSaCode.sa_code)
    if scope is not None:
        q = q.where(IncidentSaCode.district_id.in_(scope))
    rows = db.execute(q).scalars().all()
    districts = {d.id: d.name for d in db.execute(select(District)).scalars()}
    teams = {t.id: t.code for t in db.execute(select(Team).where(Team.id.in_({r.team_id for r in rows if r.team_id}))).scalars()} if any(r.team_id for r in rows) else {}
    return [SaCodeOut(id=r.id, sa_code=r.sa_code, district_id=r.district_id, district=districts.get(r.district_id), team_id=r.team_id, team_code=teams.get(r.team_id), sa_name=r.sa_name, supervisor_name=r.supervisor_name, supervisor_phone=r.supervisor_phone, active=r.active) for r in rows]


def _district_by(db: Session, text: str | None) -> District | None:
    if not text:
        return None
    t = text.strip()
    return db.execute(select(District).where((func.lower(District.name) == t.lower()) | (District.code == t))).scalars().first()


def upsert_sa_codes(db: Session, items: list[SaCodeIn]) -> SaCodeImportOut:
    created = updated = unknown = 0
    errors = []
    for i, it in enumerate(items, start=1):
        district = _district_by(db, it.district)
        team = db.execute(select(Team).where(Team.code == (it.team_code or "").strip())).scalars().first() if it.team_code else None
        if district is None and team is not None:
            district = db.get(District, team.district_id)
        if district is None:
            unknown += 1
            if it.district:
                errors.append(f"Row {i}: district '{it.district}' not found")
        row = db.execute(select(IncidentSaCode).where(IncidentSaCode.sa_code == it.sa_code)).scalars().first()
        if row is None:
            row = IncidentSaCode(sa_code=it.sa_code)
            db.add(row)
            created += 1
        else:
            updated += 1
        row.district_id = district.id if district else row.district_id
        row.team_id = team.id if team else row.team_id
        row.sa_name = (it.sa_name or "").strip() or row.sa_name
        row.supervisor_name = (it.supervisor_name or "").strip() or row.supervisor_name
        row.supervisor_phone = (it.supervisor_phone or "").strip() or row.supervisor_phone
        row.active = True
    db.flush()
    return SaCodeImportOut(created=created, updated=updated, unknown_district=unknown, errors=errors[:20])


def parse_sa_file(filename: str, data: bytes) -> list[SaCodeIn]:
    """Columns (any order, header row): sa_code, district, team_code, sa_name, supervisor_name, supervisor_phone."""
    rows: list[dict] = []
    if filename.lower().endswith((".xlsx", ".xlsm")):
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower().replace(" ", "_") for h in next(it, [])]
        for r in it:
            if any(c is not None for c in r):
                rows.append({header[i]: ("" if c is None else str(c).strip()) for i, c in enumerate(r) if i < len(header)})
    else:
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig", errors="replace")))
        rows = [{(k or "").strip().lower().replace(" ", "_"): (v or "").strip() for k, v in r.items()} for r in reader]
    out = []
    for r in rows:
        code = r.get("sa_code") or r.get("sa") or r.get("code") or ""
        code = "".join(ch for ch in code if ch.isdigit()).zfill(12) if code else ""
        if not spec.SA_CODE_RE.match(code):
            continue
        out.append(SaCodeIn(sa_code=code, district=r.get("district") or None, team_code=r.get("team_code") or r.get("sa_code_frame") or None, sa_name=r.get("sa_name") or r.get("name") or None,
                            supervisor_name=r.get("supervisor_name") or r.get("supervisor") or None, supervisor_phone=r.get("supervisor_phone") or r.get("phone") or None))
    if not out:
        raise HTTPException(400, "No SA code found: the file needs a header row with sa_code (12 digits) and district")
    return out


# ---- export ------------------------------------------------------------------------------------------


def export(db: Session, user: User, f: dict) -> report_service.Report:
    d = dashboard(db, user, f)
    rows = db.execute(_query(db, user, f).options(selectinload(Incident.updates)).order_by(Incident.reported_at.desc())).scalars().all()
    names = Names(db, rows)
    outs = [to_out(x, names, with_updates=True) for x in rows]
    headers = ["Number", "Status", "Severity", "Category", "Sub-category", "Phase", "Title", "Description", "District", "Region", "SA code", "SA / team", "EA code", "Locality", "Occurred", "Reported", "Reported by", "Reporter", "Reporter role", "Phone", "Application code",
               "Teams affected", "Staff affected", "Households affected", "Days lost", "Tablets", "Reference", "Level", "Owner", "Due", "Overdue", "Acknowledged", "Hours to acknowledge", "Resolved", "Hours to resolve", "Closed", "Resolution", "Root cause fixed", "Lessons", "Reject reason", "Action taken on the spot", "Help needed"]
    reg = [[x.number, x.status_label, x.severity_label, x.category_label, x.subcategory_label or "", x.phase_label, x.title, x.description, x.district or "", x.region or "", x.sa_code or "", x.team_code or "", x.ea_code or "", x.locality or "", x.occurred_at, x.reported_at,
            x.reported_by_name or "public page", x.reporter_name or "", x.reporter_role_label or "", x.reporter_phone or "", x.reporter_code or "", x.teams_affected, x.staff_affected, x.households_affected, x.days_lost, x.asset_tags or "", x.reference or "", spec.label("level", x.level), x.owner_name or "",
            x.due_at, "Yes" if x.overdue else "", x.acknowledged_at, x.hours_to_acknowledge, x.resolved_at, x.hours_to_resolve, x.closed_at, x.resolution or "", "" if x.root_cause_fixed is None else ("Yes" if x.root_cause_fixed else "No"), x.lessons or "", x.reject_reason or "", x.action_taken or "", x.help_needed or ""] for x in outs]
    updates = [[x.number, u.at, u.author or "", u.type_label, u.old_value or "", u.new_value or "", u.text or ""] for x in outs for u in x.updates]
    summary = [["Scope", d.scope], ["Incidents", d.total], ["Open", d.open], ["New, not acknowledged", d.new_unacknowledged], ["Critical open", d.critical_open], ["Overdue", d.overdue], ["Due today", d.due_today],
               ["Reported today", d.reported_today], ["Reported in the last 7 days", d.reported_7_days], ["Resolved in the last 7 days", d.resolved_7_days], ["Closed", d.closed], ["Rejected", d.rejected],
               ["Median hours to acknowledge", d.hours_to_acknowledge_median], ["Median hours to resolve", d.hours_to_resolve_median],
               *[[f"Severity: {c.label}", c.count] for c in d.by_severity], *[[f"Status: {c.label}", c.count] for c in d.by_status], *[[f"Phase: {c.label}", c.count] for c in d.by_phase]]
    fl = {k: v for k, v in d.filters.items() if v}
    return report_service.Report(
        kind="incidents", title="Census Incident Log", generated_at=_now(), filters="; ".join(f"{k}={v}" for k, v in fl.items()) or "none",
        sheets=[
            report_service.Sheet("Summary", ["Measure", "Value"], summary),
            report_service.Sheet("Needs action", ["Number", "Severity", "Status", "District", "Title", "Owner", "Due", "Overdue", "Reported"], [[x.number, x.severity_label, x.status_label, x.district or "", x.title, x.owner_name or "", x.due_at, "Yes" if x.overdue else "", x.reported_at] for x in d.needs_action], landscape=True),
            report_service.Sheet("By district", ["District", "Region", "Open", "New", "Critical open", "Overdue", "Resolved", "Total", "Per 100 EAs", "Median hours to resolve"], [[r.district, r.region, r.open, r.new, r.critical_open, r.overdue, r.resolved, r.total, r.per_100_eas, r.median_hours_to_resolve] for r in d.by_district]),
            report_service.Sheet("By category", ["Category", "Total", "Open", "Critical", "Major", "Minor", "Last 7 days", "Previous 7 days"], [[r.label, r.total, r.open, r.critical, r.major, r.minor, r.last_7_days, r.previous_7_days] for r in d.by_category]),
            report_service.Sheet("By owner", ["Owner", "Open", "Overdue", "Resolved"], [[r.owner, r.open, r.overdue, r.resolved] for r in d.by_owner]),
            report_service.Sheet("Trend", ["Day", "Reported", "Resolved", "Open at end of day"], [[r.day, r.reported, r.resolved, r.open_end_of_day] for r in d.trend]),
            report_service.Sheet("Register", headers, reg, landscape=True),
            report_service.Sheet("Updates", ["Number", "At", "By", "Type", "From", "To", "Text"], updates, landscape=True),
        ],
    )
